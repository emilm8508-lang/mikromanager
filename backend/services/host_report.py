"""
Export of the host-monitoring data (services/host_monitor.py) as
  - a finished Word report for non-technical readers (summary table, per-host
    findings in plain language, the log lines that back each cause),
  - CSV (the merged event timeline, one row per event — for spreadsheets),
  - JSON (everything: outages, findings, indicators, presence).

All wording comes from `labels` — the UI's own hostmon.* translations,
flattened to {"cause.port_link_down": "...", ...} and sent by the browser — so
the report is in the user's language and the texts live in one place (the
locale files) instead of being duplicated here. A missing label falls back to
its key, which is ugly but never fails the export.
"""
import csv
import io
from datetime import datetime, timedelta
from typing import Dict, List, Optional

from docx import Document
from docx.shared import Pt

from services import host_monitor

_MAX_OUTAGE_ROWS = 30
_MAX_EVIDENCE_ROWS = 25
_MAX_LABELS = 400
_MAX_LABEL_LEN = 2000


def clean_labels(raw) -> Dict[str, str]:
    """Labels arrive from the browser — keep only short str->str pairs."""
    if not isinstance(raw, dict):
        return {}
    out = {}
    for k, v in list(raw.items())[:_MAX_LABELS]:
        if isinstance(k, str) and isinstance(v, str):
            out[k] = v[:_MAX_LABEL_LEN]
    return out


def _l(labels: Dict[str, str], key: str, **params) -> str:
    text = labels.get(key, key)
    for k, v in params.items():
        text = text.replace("{{" + k + "}}", str(v))
    return text


def _dt(iso: Optional[str]) -> Optional[datetime]:
    try:
        return datetime.fromisoformat(iso) if iso else None
    except ValueError:
        return None


def fmt_time(iso: Optional[str]) -> str:
    d = _dt(iso)
    return d.strftime("%Y-%m-%d %H:%M:%S") if d else "—"


def fmt_duration(sec: Optional[float]) -> str:
    if sec is None:
        return "—"
    sec = int(sec)
    if sec < 60:
        return f"{sec} s"
    m = sec // 60
    if m < 60:
        return f"{m} min {sec % 60} s"
    h = m // 60
    return f"{h} h {m % 60} min" if h < 24 else f"{h // 24} d {h % 24} h"


def describe_event(ev: dict, labels: Dict[str, str]) -> str:
    d = ev.get("data") or {}
    kind = ev.get("kind")
    if kind == "probe_down":
        return _l(labels, "kind.probe_down", target=d.get("target", ""))
    if kind == "probe_up":
        return _l(labels, "kind.probe_up", target=d.get("target", ""), duration=fmt_duration(d.get("duration_sec")))
    if kind in ("ip_changed", "mac_changed"):
        return _l(labels, f"kind.{kind}", old=d.get("old", ""), new=d.get("new", ""))
    if kind == "port_changed":
        return _l(labels, "kind.port_changed", old=", ".join(d.get("old") or []), new=", ".join(d.get("new") or []))
    return ev.get("message") or ""


def _csv_safe(value) -> str:
    """Same OWASP formula-injection guard as api/vuln_scan.py: log lines are
    device-supplied text, and a cell starting with = + - @ would be run as a
    formula when the file is opened in a spreadsheet."""
    s = "" if value is None else str(value)
    return "'" + s if s and s[0] in ("=", "+", "-", "@", "\t", "\r") else s


def availability(report: dict, now: datetime) -> dict:
    """Share of the observed window the host was reachable. The window starts
    when monitoring of this host started if that is later than the requested
    range, so a host added yesterday isn't credited with 29 days of
    'available' time nobody measured."""
    range_start = now - timedelta(hours=report["hours"])
    created = _dt(report["host"].get("created_at")) or range_start
    win_start = max(range_start, created)
    observed = (now - win_start).total_seconds()
    down = 0.0
    for o in report["outages"]:
        s = max(_dt(o["start"]) or win_start, win_start)
        e = _dt(o["end"]) or now
        if e > s:
            down += (e - s).total_seconds()
    pct = None if observed < 60 else max(0.0, 100.0 * (1 - down / observed))
    return {"observed_since": win_start.isoformat(), "observed_sec": observed, "downtime_sec": down, "pct": pct}


def collect(host_ids: Optional[List[int]], hours: int) -> dict:
    """Reports for the selected hosts (None/empty = every monitored host)."""
    now = datetime.now()
    ids = host_ids or [h["id"] for h in host_monitor.list_hosts()]
    reports = []
    for hid in ids:
        try:
            r = host_monitor.get_report(hid, hours)
        except LookupError:
            continue
        r["availability"] = availability(r, now)
        reports.append(r)
    return {"generated_at": now.isoformat(), "hours": reports[0]["hours"] if reports else hours,
            "range_start": (now - timedelta(hours=hours)).isoformat(), "hosts": reports}


# ── CSV ──────────────────────────────────────────────────────────────────────

def render_csv(data: dict, labels: Dict[str, str]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["time", "host", "source", "device", "kind", "matched_on", "severity", "description", "topics"])
    rows = []
    for r in data["hosts"]:
        for ev in r["events"]:
            rows.append((ev["ts"], r["host"]["name"], ev))
    for ts, name, ev in sorted(rows, key=lambda x: x[0]):
        w.writerow([_csv_safe(v) for v in (
            fmt_time(ts), name, ev["source"], ev.get("device") or "", ev.get("kind") or "",
            ev.get("matched_on") or "", ev["severity"], describe_event(ev, labels), ev.get("topics") or "")])
    return buf.getvalue()


# ── Word ─────────────────────────────────────────────────────────────────────

def _table(doc, headers: List[str], rows: List[List[str]]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return t


def _cause_title(labels, code):
    return _l(labels, f"cause.{code}")


def render_docx(data: dict, labels: Dict[str, str]) -> bytes:
    L = lambda k, **p: _l(labels, k, **p)
    doc = Document()
    doc.add_heading(L("report.title"), level=1)
    doc.add_paragraph(f"{L('report.period')}: {fmt_time(data['range_start'])} – {fmt_time(data['generated_at'])}"
                      f"   |   {L('report.generated')}: {fmt_time(data['generated_at'])}")
    doc.add_paragraph(L("report.howToRead"))

    hosts = data["hosts"]
    doc.add_heading(L("report.summaryTitle"), level=2)
    if not hosts:
        doc.add_paragraph(L("report.noData"))
    else:
        rows = []
        for r in hosts:
            h, a, s = r["host"], r["availability"], r["summary"]
            rows.append([
                h["name"], f"{h.get('resolved_ip') or h.get('ip') or '—'} / {h.get('resolved_mac') or h.get('mac') or '—'}",
                f"{a['pct']:.2f} %" if a["pct"] is not None else "—",
                str(len(r["outages"])), fmt_duration(a["downtime_sec"]) if r["outages"] else "—",
                _cause_title(labels, s["code"]) if s else "—",
            ])
        _table(doc, [L("report.colHost"), L("report.colAddress"), L("report.colAvailability"),
                     L("report.colOutages"), L("report.colDowntime"), L("report.colMainCause")], rows)

    for r in hosts:
        h, a, s = r["host"], r["availability"], r["summary"]
        events_by_id = {e["id"]: e for e in r["events"]}
        doc.add_heading(h["name"], level=2)
        state = {"up": L("up"), "down": L("down")}.get(h.get("probe_state"), L("unknown"))
        doc.add_paragraph(f"{L('report.statusNow')}: {state}   |   {L('report.observedSince')}: {fmt_time(a['observed_since'])}"
                          + (f"   |   {L('report.colAvailability')}: {a['pct']:.2f} %" if a["pct"] is not None else ""))

        if not r["outages"]:
            doc.add_paragraph(L("report.noOutagesText"))
        else:
            doc.add_paragraph(L("report.outagesText", n=len(r["outages"]), total=fmt_duration(a["downtime_sec"])))
            if s:
                doc.add_heading(L("causeTitle"), level=3)
                p = doc.add_paragraph()
                p.add_run(_cause_title(labels, s["code"])).bold = True
                doc.add_paragraph(L(f"causeHint.{s['code']}"))
                doc.add_paragraph(L("causeCount", n=s["outages_with_cause"], total=s["outages"]))
                if s.get("periodic_sec"):
                    doc.add_paragraph(L("periodic", interval=fmt_duration(s["periodic_sec"])))

            doc.add_heading(L("outagesTitle"), level=3)
            findings = sorted(r["findings"], key=lambda f: f["start"], reverse=True)
            rows = []
            for f in findings[:_MAX_OUTAGE_ROWS]:
                cause = _cause_title(labels, f["code"])
                if f.get("confidence"):
                    cause += f" ({L('confidence.' + f['confidence'])})"
                rows.append([fmt_time(f["start"]), fmt_duration(f["duration_sec"]) if f["end"] else L("ongoing"), cause])
            _table(doc, [L("colStart"), L("colDuration"), L("colCause")], rows)
            if len(findings) > _MAX_OUTAGE_ROWS:
                doc.add_paragraph(f"… +{len(findings) - _MAX_OUTAGE_ROWS}")

            evidence_ids = {i for f in r["findings"] for i in f["evidence"]}
            ev = sorted((events_by_id[i] for i in evidence_ids if i in events_by_id), key=lambda e: e["ts"])
            if ev:
                doc.add_heading(L("report.relatedEvents"), level=3)
                _table(doc, [L("report.colTime"), L("report.colSource"), L("report.colEvent")],
                       [[fmt_time(e["ts"]), L(f"source.{e['source']}"),
                         (f"[{e['device']}] " if e.get("device") else "") + describe_event(e, labels)]
                        for e in ev[:_MAX_EVIDENCE_ROWS]])

        if r["indicators"]:
            doc.add_heading(L("report.indicatorsTitle"), level=3)
            for i in r["indicators"]:
                doc.add_paragraph(f"{_cause_title(labels, i['code'])} ×{i['count']} — {L('causeHint.' + i['code'])}",
                                  style="List Bullet")

        pres = (r.get("presence") or {}).get("entries") or []
        doc.add_heading(L("presenceTitle"), level=3)
        if not pres:
            doc.add_paragraph(L("presenceNone"))
        for p in pres:
            if p["type"] == "port":
                txt = f"{p['device']}: port {p.get('interface')}" + (f" — {L('uplink')}" if p.get("uplink") else "")
            elif p["type"] == "dhcp":
                txt = f"{p['device']}: DHCP {p.get('address')} {p.get('status') or ''}".strip()
            else:
                txt = f"{p['device']}: ARP {p.get('address')} → {p.get('mac')}"
            doc.add_paragraph(txt, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph(L("report.disclaimer"))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
